import io
import re
import time
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from database import (
    initialize_database,
    email_existe,
    inscrire_administrateur,
    authentifier_utilisateur,
    obtenir_utilisateur,
    changer_mot_de_passe,
    ajouter_transaction,
    lister_transactions,
    obtenir_transaction,
    modifier_transaction,
    supprimer_transaction,
    calculer_total,
    obtenir_solde_initial,
    modifier_solde_initial,
)
from rapports import page_rapports_avances

# ============================================================
# CONFIGURATION
# ============================================================

st.set_page_config(page_title="FinTech PME", page_icon="💼", layout="wide")

initialize_database()

MOYENS_PAIEMENT = [
    "Espèces", "Mobile Money", "Virement bancaire",
    "Carte bancaire", "Chèque", "Autre",
]

MOT_DE_PASSE_MIN = 8
MAX_TENTATIVES = 5
DUREE_BLOCAGE_SEC = 300

CONFIG = {
    "recettes": {
        "titre": "💰 Gestion des recettes",
        "un": "une recette",
        "nouveau_titre": "Nouvelle recette",
        "tiers": "Source",
        "placeholder": "Vente, prestation, etc.",
        "ok_ajout": "Recette enregistrée avec succès.",
        "ok_modif": "Recette modifiée avec succès.",
        "ok_suppr": "Recette supprimée.",
        "export": "rapport_recettes.xlsx",
        "feuille": "Recettes",
        "rapport": "Rapport des recettes",
    },
    "depenses": {
        "titre": "💸 Gestion des dépenses",
        "un": "une dépense",
        "nouveau_titre": "Nouvelle dépense",
        "tiers": "Bénéficiaire",
        "placeholder": "Loyer, salaire, achat, etc.",
        "ok_ajout": "Dépense enregistrée avec succès.",
        "ok_modif": "Dépense modifiée avec succès.",
        "ok_suppr": "Dépense supprimée.",
        "export": "rapport_depenses.xlsx",
        "feuille": "Depenses",
        "rapport": "Rapport des dépenses",
    },
}

# ============================================================
# SESSION
# ============================================================

DEFAUTS_SESSION = {
    "connecte": False,
    "utilisateur_id": None,
    "entreprise_id": None,
    "nom_utilisateur": "",
    "email_utilisateur": "",
    "tentatives": 0,
    "bloque_jusqua": 0.0,
}

for cle, valeur in DEFAUTS_SESSION.items():
    if cle not in st.session_state:
        st.session_state[cle] = valeur


def deconnecter():
    for cle, valeur in DEFAUTS_SESSION.items():
        st.session_state[cle] = valeur


def flash(message):
    """Message affiché (toast) après le prochain rechargement."""
    st.session_state["flash"] = message


if "flash" in st.session_state:
    st.toast(st.session_state.pop("flash"), icon="✅")


# ============================================================
# UTILITAIRES
# ============================================================

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def format_fcfa(montant):
    return f"{montant:,.0f} FCFA".replace(",", " ")


def nettoyer_texte(valeur):
    return "" if valeur is None else str(valeur)


def charger_df(table, entreprise_id):
    lignes = lister_transactions(table, entreprise_id)
    return pd.DataFrame(
        lignes,
        columns=["ID", "Date", "Montant", "Catégorie", "Tiers",
                 "Paiement", "Description"],
    )


def situation(entreprise_id):
    solde_initial = obtenir_solde_initial(entreprise_id)
    total_recettes = calculer_total("recettes", entreprise_id)
    total_depenses = calculer_total("depenses", entreprise_id)
    resultat = total_recettes - total_depenses
    solde_actuel = solde_initial + resultat
    return solde_initial, total_recettes, total_depenses, resultat, solde_actuel


def exporter_excel(df, nom_feuille, titre, administrateur):
    buffer = io.BytesIO()
    ligne_entete = 5  # startrow=4 -> en-têtes en ligne 5

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=nom_feuille, startrow=4)
        feuille = writer.sheets[nom_feuille]

        feuille["A1"] = "FINTECH PME"
        feuille["A2"] = titre
        feuille["A3"] = f"Administrateur : {administrateur}"

        derniere_ligne = ligne_entete + len(df)

        for index, colonne in enumerate(df.columns, start=1):
            lettre = get_column_letter(index)
            longueur = len(str(colonne))
            for ligne in range(ligne_entete + 1, derniere_ligne + 1):
                valeur = feuille[f"{lettre}{ligne}"].value
                if valeur is not None:
                    longueur = max(longueur, len(str(valeur)))
            feuille.column_dimensions[lettre].width = min(longueur + 3, 40)

        feuille.freeze_panes = f"A{ligne_entete + 1}"
        feuille.auto_filter.ref = (
            f"A{ligne_entete}:{get_column_letter(len(df.columns))}{derniere_ligne}"
        )

    return buffer.getvalue()


# ============================================================
# AUTHENTIFICATION
# ============================================================

def page_authentification():

    st.title("💼 FinTech PME")
    st.markdown(
        "### Plateforme de gestion financière pour les petites et moyennes entreprises"
    )
    st.divider()

    onglet_connexion, onglet_inscription = st.tabs(
        ["🔐 Connexion administrateur", "📝 Créer un compte"]
    )

    # ---------------- Connexion ----------------
    with onglet_connexion:

        st.subheader("Connexion administrateur")

        with st.form("form_connexion"):
            email = st.text_input("Adresse e-mail", placeholder="admin@entreprise.com")
            mot_de_passe = st.text_input("Mot de passe", type="password")
            connexion = st.form_submit_button("Se connecter", use_container_width=True)

        if connexion:
            maintenant = time.time()

            if maintenant < st.session_state.bloque_jusqua:
                reste = int(st.session_state.bloque_jusqua - maintenant)
                st.error(f"Trop de tentatives. Réessayez dans {reste} secondes.")

            elif not email or not mot_de_passe:
                st.error("Veuillez renseigner votre adresse e-mail et votre mot de passe.")

            else:
                utilisateur = authentifier_utilisateur(email, mot_de_passe)

                if utilisateur:
                    st.session_state.connecte = True
                    st.session_state.utilisateur_id = utilisateur[0]
                    st.session_state.nom_utilisateur = utilisateur[1]
                    st.session_state.email_utilisateur = utilisateur[2]
                    st.session_state.entreprise_id = utilisateur[3]
                    st.session_state.tentatives = 0
                    flash("Connexion réussie.")
                    st.rerun()
                else:
                    st.session_state.tentatives += 1
                    if st.session_state.tentatives >= MAX_TENTATIVES:
                        st.session_state.bloque_jusqua = maintenant + DUREE_BLOCAGE_SEC
                        st.session_state.tentatives = 0
                        st.error("Trop de tentatives. Compte bloqué pendant 5 minutes.")
                    else:
                        st.error("Adresse e-mail ou mot de passe incorrect.")

    # ---------------- Inscription ----------------
    with onglet_inscription:

        st.subheader("Créer un compte administrateur")
        st.info(
            "L'inscription crée automatiquement votre entreprise "
            "et votre compte administrateur."
        )

        with st.form("form_inscription"):
            st.markdown("#### Informations de l'administrateur")
            nom_admin = st.text_input("Nom complet")
            email_admin = st.text_input("Adresse e-mail")
            mot_de_passe_admin = st.text_input("Mot de passe", type="password")
            confirmation = st.text_input("Confirmer le mot de passe", type="password")

            st.markdown("#### Informations de l'entreprise")
            nom_entreprise = st.text_input("Nom de l'entreprise")
            secteur = st.text_input(
                "Secteur d'activité",
                placeholder="Commerce, restauration, services...",
            )
            adresse = st.text_input("Adresse")
            telephone = st.text_input("Téléphone")
            email_entreprise = st.text_input("E-mail de l'entreprise")

            inscription = st.form_submit_button("Créer mon compte", use_container_width=True)

        if inscription:

            if not nom_admin.strip():
                st.error("Veuillez renseigner votre nom.")
            elif not EMAIL_REGEX.match(email_admin.strip()):
                st.error("Veuillez renseigner une adresse e-mail valide.")
            elif len(mot_de_passe_admin) < MOT_DE_PASSE_MIN:
                st.error(
                    f"Le mot de passe doit contenir au moins {MOT_DE_PASSE_MIN} caractères."
                )
            elif mot_de_passe_admin != confirmation:
                st.error("Les deux mots de passe ne correspondent pas.")
            elif not nom_entreprise.strip():
                st.error("Veuillez renseigner le nom de l'entreprise.")
            elif email_entreprise.strip() and not EMAIL_REGEX.match(email_entreprise.strip()):
                st.error("L'adresse e-mail de l'entreprise n'est pas valide.")
            elif email_existe(email_admin):
                st.error("Cette adresse e-mail possède déjà un compte.")
            else:
                try:
                    inscrire_administrateur(
                        nom_admin.strip(),
                        email_admin,
                        mot_de_passe_admin,
                        nom_entreprise.strip(),
                        secteur.strip(),
                        adresse.strip(),
                        telephone.strip(),
                        email_entreprise.strip(),
                    )
                    st.success("Compte administrateur créé avec succès.")
                    st.info(
                        "Vous pouvez maintenant vous connecter "
                        "avec votre adresse e-mail et votre mot de passe."
                    )
                except ValueError:
                    st.error("Cette adresse e-mail possède déjà un compte.")
                except Exception:
                    st.error("Une erreur est survenue. Veuillez réessayer.")


if not st.session_state.connecte:
    page_authentification()
    st.stop()


# ============================================================
# VÉRIFICATION DE LA SESSION
# ============================================================

utilisateur_actuel = obtenir_utilisateur(st.session_state.utilisateur_id)

if utilisateur_actuel is None or utilisateur_actuel[5] != 1:
    deconnecter()
    st.rerun()

if utilisateur_actuel[4] != "admin":
    st.error("Accès refusé. Seuls les administrateurs peuvent utiliser la plateforme.")
    st.stop()

entreprise_id = st.session_state.entreprise_id


# ============================================================
# PAGES
# ============================================================

def page_tableau_de_bord(eid):

    st.title("📊 Tableau de bord")
    st.write("Bienvenue dans votre espace de gestion financière.")
    st.divider()

    solde_initial, total_recettes, total_depenses, resultat, solde_actuel = situation(eid)

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Solde actuel", format_fcfa(solde_actuel))
    c2.metric("Total recettes", format_fcfa(total_recettes))
    c3.metric("Total dépenses", format_fcfa(total_depenses))
    c4.metric("Résultat", format_fcfa(resultat))

    st.divider()

    df_rec = charger_df("recettes", eid).assign(Type="Recettes")
    df_dep = charger_df("depenses", eid).assign(Type="Dépenses")

    st.subheader("📌 Indicateurs financiers")

    marge = (resultat / total_recettes * 100) if total_recettes > 0 else 0
    moyenne = (total_recettes / len(df_rec)) if len(df_rec) > 0 else 0

    c1, c2, c3 = st.columns(3)
    c1.metric("Nombre de transactions", len(df_rec) + len(df_dep))
    c2.metric("Marge financière", f"{marge:.1f} %")
    c3.metric("Recette moyenne", format_fcfa(moyenne))

    st.divider()

    tous = pd.concat([df_rec, df_dep], ignore_index=True)

    st.subheader("📈 Évolution financière")

    if tous.empty:
        st.info("Aucune donnée financière disponible pour afficher le graphique.")
    else:
        graphique = (
            tous.assign(Mois=tous["Date"].str[:7])
            .pivot_table(
                index="Mois", columns="Type", values="Montant",
                aggfunc="sum", fill_value=0,
            )
            .sort_index()
        )
        st.line_chart(graphique)

    st.divider()

    st.subheader("🧾 Dernières transactions")

    if tous.empty:
        st.info("Aucune transaction enregistrée.")
    else:
        dernieres = (
            tous.sort_values(["Date", "ID"], ascending=False)
            .head(10)[["Date", "Type", "Catégorie", "Montant", "Description"]]
        )
        st.dataframe(dernieres, use_container_width=True, hide_index=True)


def page_transactions(table, eid):

    cfg = CONFIG[table]
    libelle_tiers = cfg["tiers"]

    st.title(cfg["titre"])

    onglet_ajout, onglet_liste = st.tabs(
        [f"➕ Ajouter {cfg['un']}", f"📋 Gérer les {table}"]
    )

    # ---------------- Ajout ----------------
    with onglet_ajout:

        st.subheader(cfg["nouveau_titre"])

        with st.form(f"form_ajout_{table}"):
            date_op = st.date_input("Date", value=date.today())
            montant = st.number_input("Montant (FCFA)", min_value=0.0, step=100.0)
            categorie = st.text_input("Catégorie", placeholder=cfg["placeholder"])
            tiers = st.text_input(libelle_tiers)
            moyen = st.selectbox("Moyen de paiement", MOYENS_PAIEMENT)
            description = st.text_area("Description")
            enregistrer = st.form_submit_button(
                f"Enregistrer {cfg['un'].replace('une', 'la')}",
                use_container_width=True,
            )

        if enregistrer:
            if montant <= 0:
                st.error("Le montant doit être supérieur à zéro.")
            elif not categorie.strip():
                st.error("Veuillez renseigner la catégorie.")
            else:
                ajouter_transaction(
                    table, eid, date_op.strftime("%Y-%m-%d"), montant,
                    categorie.strip(), tiers.strip(), moyen, description.strip(),
                )
                flash(cfg["ok_ajout"])
                st.rerun()

    # ---------------- Gestion ----------------
    with onglet_liste:

        df = charger_df(table, eid)

        if df.empty:
            st.info(f"Aucune {cfg['un'].split()[-1]} enregistrée.")
            return

        recherche = st.text_input(
            "🔎 Rechercher",
            placeholder=f"Catégorie, {libelle_tiers.lower()} ou description...",
            key=f"recherche_{table}",
        )

        categories = sorted(df["Catégorie"].dropna().unique().tolist())
        filtre = st.selectbox(
            "Filtrer par catégorie", ["Toutes"] + categories, key=f"filtre_{table}"
        )

        texte = (
            df["Catégorie"].fillna("") + " "
            + df["Tiers"].fillna("") + " "
            + df["Description"].fillna("")
        ).str.lower()

        masque = texte.str.contains(recherche.lower().strip(), regex=False)
        if filtre != "Toutes":
            masque &= df["Catégorie"] == filtre

        df_f = df[masque]

        if df_f.empty:
            st.info("Aucune ligne ne correspond aux critères.")
            return

        df_affiche = df_f.rename(columns={"Tiers": libelle_tiers})
        st.dataframe(df_affiche, use_container_width=True, hide_index=True)

        # ----- Modifier / supprimer -----
        st.divider()
        st.subheader(f"Modifier ou supprimer {cfg['un']}")

        ids = [int(i) for i in df_f["ID"]]
        etiquettes = {
            int(r.ID): f"#{r.ID} · {r.Date} · {r.Catégorie} · {format_fcfa(r.Montant)}"
            for r in df_f.itertuples()
        }

        selection = st.selectbox(
            "Sélectionner une ligne",
            ids,
            format_func=lambda i: etiquettes[i],
            key=f"selection_{table}",
        )

        ligne = obtenir_transaction(table, selection, eid)

        if ligne:
            k = f"{table}_{selection}"
            col1, col2 = st.columns(2)

            with col1:
                n_date = st.date_input(
                    "Date",
                    value=pd.to_datetime(ligne[1]).date(),
                    key=f"date_{k}",
                )
                n_montant = st.number_input(
                    "Montant", value=float(ligne[2]), min_value=0.0,
                    step=100.0, key=f"montant_{k}",
                )
                n_categorie = st.text_input(
                    "Catégorie", value=nettoyer_texte(ligne[3]), key=f"cat_{k}"
                )

            with col2:
                n_tiers = st.text_input(
                    libelle_tiers, value=nettoyer_texte(ligne[4]), key=f"tiers_{k}"
                )
                n_moyen = st.selectbox(
                    "Moyen de paiement",
                    MOYENS_PAIEMENT,
                    index=MOYENS_PAIEMENT.index(ligne[5])
                    if ligne[5] in MOYENS_PAIEMENT else 0,
                    key=f"moyen_{k}",
                )
                n_description = st.text_area(
                    "Description", value=nettoyer_texte(ligne[6]), key=f"desc_{k}"
                )

            col_modif, col_suppr = st.columns(2)

            with col_modif:
                if st.button(
                    "💾 Enregistrer les modifications",
                    use_container_width=True,
                    key=f"btn_modif_{k}",
                ):
                    if n_montant <= 0:
                        st.error("Le montant doit être supérieur à zéro.")
                    elif not n_categorie.strip():
                        st.error("Veuillez renseigner la catégorie.")
                    else:
                        modifie = modifier_transaction(
                            table, selection, eid,
                            n_date.strftime("%Y-%m-%d"), n_montant,
                            n_categorie.strip(), n_tiers.strip(),
                            n_moyen, n_description.strip(),
                        )
                        if modifie:
                            flash(cfg["ok_modif"])
                            st.rerun()
                        else:
                            st.error("Modification impossible : ligne introuvable.")

            with col_suppr:
                confirme = st.checkbox(
                    "Je confirme la suppression", key=f"conf_{k}"
                )
                if st.button(
                    "🗑️ Supprimer",
                    use_container_width=True,
                    disabled=not confirme,
                    key=f"btn_suppr_{k}",
                ):
                    if supprimer_transaction(table, selection, eid):
                        flash(cfg["ok_suppr"])
                        st.rerun()
                    else:
                        st.error("Suppression impossible : ligne introuvable.")

        # ----- Export Excel -----
        st.divider()
        st.subheader("📥 Export Excel")

        st.download_button(
            f"📥 Télécharger les {table} en Excel",
            data=exporter_excel(
                df_affiche, cfg["feuille"], cfg["rapport"],
                st.session_state.nom_utilisateur,
            ),
            file_name=cfg["export"],
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )


def page_tresorerie(eid):

    st.title("🏦 Gestion de la trésorerie")
    st.subheader("Situation financière")

    solde_initial, total_recettes, total_depenses, resultat, solde_actuel = situation(eid)

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Solde initial", format_fcfa(solde_initial))
    c2.metric("Total recettes", format_fcfa(total_recettes))
    c3.metric("Total dépenses", format_fcfa(total_depenses))
    c4.metric("Solde actuel", format_fcfa(solde_actuel))

    st.divider()
    st.subheader("Modifier le solde initial")

    with st.form("form_solde_initial"):
        nouveau_solde = st.number_input(
            "Solde initial (FCFA) — peut être négatif en cas de découvert",
            value=float(solde_initial),
            step=1000.0,
        )
        modifier = st.form_submit_button("💾 Enregistrer", use_container_width=True)

    if modifier:
        modifier_solde_initial(eid, nouveau_solde)
        flash("Solde initial mis à jour.")
        st.rerun()

    st.divider()
    st.subheader("Résumé")

    st.write(f"Solde initial : **{format_fcfa(solde_initial)}**")
    st.write(f"Total des recettes : **{format_fcfa(total_recettes)}**")
    st.write(f"Total des dépenses : **{format_fcfa(total_depenses)}**")
    st.write(f"Résultat : **{format_fcfa(resultat)}**")
    st.write(f"Solde actuel : **{format_fcfa(solde_actuel)}**")


def page_rapports(eid):

    st.title("📈 Rapports financiers")

    st.info(
        "Le module de rapports avancés sera développé dans la prochaine étape."
    )

    st.subheader("Résumé financier")

    solde_initial, total_recettes, total_depenses, resultat, solde_actuel = situation(eid)

    st.dataframe(
        pd.DataFrame({
            "Indicateur": [
                "Solde initial", "Total recettes", "Total dépenses",
                "Résultat", "Solde actuel",
            ],
            "Montant (FCFA)": [
                solde_initial, total_recettes, total_depenses,
                resultat, solde_actuel,
            ],
        }),
        use_container_width=True,
        hide_index=True,
    )


def page_parametres():

    st.title("⚙️ Paramètres")

    st.subheader("👤 Administrateur")
    st.write(f"**Nom :** {st.session_state.nom_utilisateur}")
    st.write(f"**E-mail :** {st.session_state.email_utilisateur}")
    st.write("**Rôle :** Administrateur")

    st.divider()
    st.subheader("🔐 Sécurité")

    with st.form("form_mot_de_passe", clear_on_submit=True):
        ancien = st.text_input("Mot de passe actuel", type="password")
        nouveau = st.text_input("Nouveau mot de passe", type="password")
        confirmation = st.text_input("Confirmer le nouveau mot de passe", type="password")
        valider = st.form_submit_button("Changer le mot de passe", use_container_width=True)

    if valider:
        if not ancien or not nouveau:
            st.error("Veuillez remplir tous les champs.")
        elif len(nouveau) < MOT_DE_PASSE_MIN:
            st.error(
                f"Le mot de passe doit contenir au moins {MOT_DE_PASSE_MIN} caractères."
            )
        elif nouveau != confirmation:
            st.error("Les deux mots de passe ne correspondent pas.")
        elif changer_mot_de_passe(st.session_state.utilisateur_id, ancien, nouveau):
            st.success("Mot de passe modifié avec succès.")
        else:
            st.error("Le mot de passe actuel est incorrect.")

    st.divider()
    st.subheader("🏢 Entreprise")
    st.info(
        "Les informations détaillées de l'entreprise seront "
        "configurables dans une prochaine version."
    )


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.title("💼 FinTech PME")
st.sidebar.write(f"**Administrateur :** {st.session_state.nom_utilisateur}")
st.sidebar.write(f"**E-mail :** {st.session_state.email_utilisateur}")
st.sidebar.divider()

menu = st.sidebar.radio(
    "Navigation",
    [
        "📊 Tableau de bord",
        "💰 Recettes",
        "💸 Dépenses",
        "🏦 Trésorerie",
        "📈 Rapports",
        "⚙️ Paramètres",
    ],
)

st.sidebar.divider()

if st.sidebar.button("🚪 Se déconnecter", use_container_width=True):
    deconnecter()
    st.rerun()


# ============================================================
# ROUTAGE
# ============================================================

if menu == "📊 Tableau de bord":
    page_tableau_de_bord(entreprise_id)
elif menu == "💰 Recettes":
    page_transactions("recettes", entreprise_id)
elif menu == "💸 Dépenses":
    page_transactions("depenses", entreprise_id)
elif menu == "🏦 Trésorerie":
    page_tresorerie(entreprise_id)
elif menu == "📈 Rapports":
    page_rapports_avances(entreprise_id)
elif menu == "⚙️ Paramètres":
    page_parametres()
